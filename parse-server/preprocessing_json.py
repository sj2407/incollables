#!/usr/bin/env python3
"""
preprocessing_json.py — JSON output wrapper for preprocessing.py.

Called by /api/parse-word to parse raw .docx files server-side.
Outputs ImportParseResult-compatible JSON to stdout.

Usage:
    python3 preprocessing_json.py <path-to-docx>

Pipeline:
  1. python-docx for structure (questions, propositions, answers, chapters, bold detection)
  2. mammoth for HTML (explanation formatting, embedded images as base64)
  3. Combines both into ImportParseResult JSON
"""

import base64
import io
import json
import re
import sys
import tempfile
from pathlib import Path

try:
    from PIL import Image as PILImage
    _HAS_PILLOW = True
except ImportError:
    _HAS_PILLOW = False

# Import from the existing preprocessing script (same directory)
sys.path.insert(0, str(Path(__file__).parent))
from preprocessing import parse_document, standardize_questions  # noqa: E402

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

import mammoth  # noqa: E402


# ---------------------------------------------------------------------------
# Chapter title cleaning
# ---------------------------------------------------------------------------

def _clean_chapter_title(title: str) -> str:
    """Strip 'Chapitre :' / 'Sous chapitre :' prefixes and non-breaking spaces."""
    t = title.replace('\xa0', ' ').strip()
    # Remove "Chapitre :" or "Sous chapitre :" prefix (colon optional)
    t = re.sub(r'^(Sous\s+)?Chapitre\s*:?\s*', '', t, flags=re.IGNORECASE).strip()
    return t if t else title.replace('\xa0', ' ').strip()


# ---------------------------------------------------------------------------
# Word image crop support
# ---------------------------------------------------------------------------

def _apply_word_crop(image_blob: bytes, blip_element, content_type: str) -> bytes:
    """
    Apply Word's a:srcRect crop metadata to an image blob if present.
    Returns the original blob unchanged when no crop is needed or Pillow is unavailable.
    """
    if not _HAS_PILLOW:
        return image_blob

    parent = blip_element.getparent()
    if parent is None:
        return image_blob

    src_rect = parent.find(qn('a:srcRect'))
    if src_rect is None:
        return image_blob

    attrs = src_rect.attrib
    # Values are in 1/1000ths of a percent (e.g. 13554 = 13.554%)
    l = int(attrs.get('l', 0))
    t = int(attrs.get('t', 0))
    r = int(attrs.get('r', 0))
    b = int(attrs.get('b', 0))

    # No actual crop — return original
    if l == 0 and t == 0 and r == 0 and b == 0:
        return image_blob

    # Invalid crop (would produce zero-dimension image) — skip
    if l + r >= 100000 or t + b >= 100000:
        return image_blob

    try:
        img = PILImage.open(io.BytesIO(image_blob))
        w, h = img.size
        box = (
            int(w * l / 100000),
            int(h * t / 100000),
            w - int(w * r / 100000),
            h - int(h * b / 100000),
        )
        cropped = img.crop(box)
        buf = io.BytesIO()
        fmt = img.format or ('PNG' if 'png' in content_type.lower() else 'JPEG')
        # JPEG doesn't support alpha — convert RGBA to RGB
        if fmt == 'JPEG' and cropped.mode == 'RGBA':
            cropped = cropped.convert('RGB')
        cropped.save(buf, format=fmt)
        return buf.getvalue()
    except Exception:
        # On any error, return the original uncropped image
        return image_blob


def _precrop_docx_images(doc) -> None:
    """
    Walk all a:blip elements in the document. For any with a:srcRect crop
    metadata, replace the image part blob with the cropped version in-memory.
    This ensures mammoth picks up cropped images when converting to HTML.

    Skips images where the same rId has conflicting crop rects (reused image
    with different crops on different instances).
    """
    if not _HAS_PILLOW:
        return

    # First pass: collect crop info per rId, detect conflicts
    crop_by_rid: dict[str, tuple] = {}
    conflict_rids: set[str] = set()

    for para in doc.paragraphs:
        for blip in para._element.findall('.//' + qn('a:blip')):
            rid = blip.get(qn('r:embed'))
            if not rid:
                continue
            parent = blip.getparent()
            if parent is None:
                continue
            src_rect = parent.find(qn('a:srcRect'))
            crop_key = tuple(sorted(src_rect.attrib.items())) if src_rect is not None else ()
            if rid in crop_by_rid and crop_by_rid[rid] != crop_key:
                conflict_rids.add(rid)
            crop_by_rid[rid] = crop_key

    # Second pass: apply crops (skip conflicts)
    done_rids: set[str] = set()
    for para in doc.paragraphs:
        for blip in para._element.findall('.//' + qn('a:blip')):
            rid = blip.get(qn('r:embed'))
            if not rid or rid in done_rids or rid in conflict_rids:
                continue
            if rid not in doc.part.rels:
                continue
            rel = doc.part.rels[rid]
            try:
                original_blob = rel.target_part.blob
                content_type = rel.target_part.content_type or 'image/png'
                cropped_blob = _apply_word_crop(original_blob, blip, content_type)
                if cropped_blob is not original_blob:
                    # Mutate the in-memory blob so mammoth sees the cropped version
                    rel.target_part._blob = cropped_blob
                    done_rids.add(rid)
            except Exception:
                pass


# ---------------------------------------------------------------------------
# Image extraction from python-docx (Word XML relationships)
# ---------------------------------------------------------------------------

def extract_stem_images(doc, question) -> list:
    """
    Extract base64-encoded images from a question's stem paragraphs.
    Uses python-docx relationship parts to get image binary data.
    """
    images = []
    sort_order = 0

    for para_idx in question.stem_para_indices:
        if para_idx >= len(doc.paragraphs):
            continue
        para = doc.paragraphs[para_idx]
        # Find all a:blip elements (embedded images)
        blips = para._element.findall('.//' + qn('a:blip'))
        for blip in blips:
            r_embed = blip.get(qn('r:embed'))
            if not r_embed:
                continue
            # Resolve the relationship to get the image part
            if r_embed not in doc.part.rels:
                continue
            rel = doc.part.rels[r_embed]
            try:
                image_blob = rel.target_part.blob
                content_type = rel.target_part.content_type or 'image/png'
                b64 = base64.b64encode(image_blob).decode('ascii')
                images.append({
                    'base64': b64,
                    'contentType': content_type,
                    'sortOrder': sort_order,
                })
                sort_order += 1
            except Exception:
                # Skip images that can't be extracted
                pass

    return images


# ---------------------------------------------------------------------------
# Explanation HTML via mammoth
# ---------------------------------------------------------------------------

def _convert_image(image):
    """Convert a mammoth image to a base64 data URI."""
    with image.open() as img_file:
        img_bytes = img_file.read()
    content_type = image.content_type or 'image/png'
    b64 = base64.b64encode(img_bytes).decode('ascii')
    return {'src': f'data:{content_type};base64,{b64}'}


def _convert_mammoth(docx_path: str) -> list:
    """
    Convert a .docx to HTML via mammoth and split into top-level elements.
    Returns list of HTML element strings.
    """
    with open(docx_path, 'rb') as f:
        mammoth_result = mammoth.convert_to_html(
            f,
            convert_image=mammoth.images.img_element(_convert_image)
        )

    html = mammoth_result.value
    if not html:
        return []

    # Split HTML into top-level block elements
    element_pattern = re.compile(
        r'(<(?:h[1-6]|p|ol|ul|table|div|blockquote)[^>]*>.*?'
        r'</(?:h[1-6]|p|ol|ul|table|div|blockquote)>)',
        re.DOTALL,
    )
    return element_pattern.findall(html)


def _match_para_to_html(para_text: str, html_elements: list, used: set) -> str | None:
    """
    Find the mammoth HTML element that matches a python-docx paragraph's text.
    Returns the HTML string, or None if no match.
    Marks used indices to avoid double-matching.
    """
    clean = para_text[:60].strip()
    if not clean:
        return None
    for idx, html_el in enumerate(html_elements):
        if idx in used:
            continue
        el_text = re.sub(r'<[^>]+>', '', html_el).strip()
        if clean[:40] in el_text:
            used.add(idx)
            return html_el
    return None


def _find_image_html_near(last_matched_idx: int, html_elements: list, used: set) -> list:
    """
    After matching a text paragraph at `last_matched_idx`, look at the next
    few HTML elements for image-only elements (contain <img but no meaningful text).
    This captures image-only explanation paragraphs that have no text to match.
    """
    found = []
    # Check the next few elements after the last matched one
    start = last_matched_idx + 1 if last_matched_idx >= 0 else 0
    for idx in range(start, min(start + 15, len(html_elements))):
        if idx in used:
            continue
        el_text = re.sub(r'<[^>]+>', '', html_elements[idx]).strip()
        has_img = '<img ' in html_elements[idx]
        if has_img and len(el_text) < 10:
            # Image-only element — grab it
            found.append(html_elements[idx])
            used.add(idx)
        elif el_text:
            # Hit a text element — stop looking
            break
    return found


def _build_html_map(
    questions: list,
    doc: Document,
    html_elements: list,
    para_attr: str,
) -> dict:
    """
    Generic builder: maps question paragraph indices (stem or explanation)
    to mammoth HTML, including image-only paragraphs.

    para_attr: 'stem_para_indices' or 'explanation_para_indices'
    Returns dict: sequential_number -> HTML string
    """
    from html import escape

    total_paras = len(doc.paragraphs)
    result_map = {}
    used = set()  # track which html_elements have been consumed

    for q in questions:
        para_indices = getattr(q, para_attr, [])
        if not para_indices:
            result_map[q.sequential_number] = ''
            continue

        # Collect paragraph info: text + whether it has images
        para_info = []
        for pi in para_indices:
            if pi >= total_paras:
                continue
            para = doc.paragraphs[pi]
            text = para.text.strip()
            has_image = bool(para._element.findall('.//' + qn('a:blip')))
            para_info.append((text, has_image))

        if not para_info:
            result_map[q.sequential_number] = ''
            continue

        matched_html_parts = []
        last_matched_idx = -1
        has_any_image = any(img for _, img in para_info)

        for text, has_image in para_info:
            # Skip the "Explication :" marker but use it to anchor position
            if re.match(r'^Explication\s*:?\s*$', text, re.IGNORECASE):
                # Find this marker in HTML to anchor for nearby image search
                anchor = _match_para_to_html(text, html_elements, used)
                if anchor:
                    last_matched_idx = html_elements.index(anchor)
                    # Don't add the marker itself to output
                continue

            if text:
                # Try to match by text content
                match = _match_para_to_html(text, html_elements, used)
                if match:
                    matched_html_parts.append(match)
                    last_matched_idx = html_elements.index(match)
                    # After a text match, check for trailing image-only elements
                    img_els = _find_image_html_near(last_matched_idx, html_elements, used)
                    matched_html_parts.extend(img_els)
                else:
                    # No HTML match — use plain text wrapped in <p>
                    matched_html_parts.append(f'<p>{escape(text)}</p>')
            elif has_image:
                # Image-only paragraph (no text) — look for <img> in HTML
                img_els = _find_image_html_near(last_matched_idx, html_elements, used)
                matched_html_parts.extend(img_els)

        # If we found no text matches but there are images, do a broader scan
        if not matched_html_parts and has_any_image and last_matched_idx >= 0:
            img_els = _find_image_html_near(last_matched_idx, html_elements, used)
            matched_html_parts.extend(img_els)

        result_map[q.sequential_number] = '\n'.join(matched_html_parts)

    return result_map


def build_explanation_html_map(docx_path: str, questions: list, doc: Document,
                                html_elements: list) -> dict:
    """Build explanation HTML map using mammoth output."""
    return _build_html_map(questions, doc, html_elements, 'explanation_para_indices')


def build_stem_html_map(docx_path: str, questions: list, doc: Document,
                         html_elements: list) -> dict:
    """Build stem HTML map using mammoth output (includes inline images)."""
    return _build_html_map(questions, doc, html_elements, 'stem_para_indices')


# ---------------------------------------------------------------------------
# Topic tree builder (mirrors TypeScript buildTopicTree)
# ---------------------------------------------------------------------------

def build_topic_tree(questions: list) -> list:
    """
    Build a nested topic tree from questions' chapterPath arrays.
    Returns a list of ParsedTopic dicts.
    """
    roots = []
    seen = {}  # path_key -> topic dict
    sort_counter = 0

    for q_data in questions:
        chapter_path = q_data.get('chapterPath', [])
        if not chapter_path:
            continue
        for depth in range(len(chapter_path)):
            path_key = '///'.join(chapter_path[:depth + 1])
            if path_key in seen:
                continue
            node = {
                'title': chapter_path[depth],
                'level': depth + 1,
                'sortOrder': sort_counter,
                'children': [],
            }
            sort_counter += 1
            seen[path_key] = node
            if depth == 0:
                roots.append(node)
            else:
                parent_key = '///'.join(chapter_path[:depth])
                parent = seen.get(parent_key)
                if parent:
                    parent['children'].append(node)

    return roots


def count_topics(nodes: list) -> int:
    """Count total topics in a nested tree."""
    total = 0
    for n in nodes:
        total += 1 + count_topics(n.get('children', []))
    return total


# ---------------------------------------------------------------------------
# Core parsing function (callable from Flask app)
# ---------------------------------------------------------------------------

def parse_docx_to_json(docx_path: str) -> dict:
    """Parse a .docx and return ImportParseResult-compatible dict."""
    doc = Document(docx_path)
    questions, chapters, stats = parse_document(doc)
    questions = standardize_questions(questions)
    _precrop_docx_images(doc)

    # Save cropped doc to temp file for mammoth
    cropped_docx_path = docx_path
    if _HAS_PILLOW:
        tmp = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        tmp.close()
        doc.save(tmp.name)
        cropped_docx_path = tmp.name

    html_elements = _convert_mammoth(cropped_docx_path)

    if cropped_docx_path != docx_path:
        Path(cropped_docx_path).unlink(missing_ok=True)

    stem_html_map = build_stem_html_map(docx_path, questions, doc, html_elements)
    explanation_map = build_explanation_html_map(docx_path, questions, doc, html_elements)

    questions_json = []
    for q in questions:
        chapter_path = [_clean_chapter_title(c) for c in q.chapter_path]
        stem_images = extract_stem_images(doc, q)
        stem_html = stem_html_map.get(q.sequential_number, '')
        expl_html = explanation_map.get(q.sequential_number, q.explanation_text or '')
        stem_text = stem_html if stem_html else q.stem_text

        questions_json.append({
            'sequentialNumber': q.sequential_number,
            'stemText': stem_text,
            'propositions': [{'letter': p.letter, 'text': p.text} for p in q.propositions],
            'answerLetters': q.answer_letters,
            'answerMissing': len(q.answer_letters) == 0,
            'explanationHtml': expl_html,
            'chapterPath': chapter_path,
            'stemImages': stem_images,
        })

    topics = build_topic_tree(questions_json)

    active_count = sum(1 for q in questions_json
        if not q['answerMissing'] and len(q['propositions']) >= 3 and len(q['answerLetters']) > 0)

    result_stats = {
        'totalQuestions': len(questions_json),
        'activeQuestions': active_count,
        'inactiveQuestions': len(questions_json) - active_count,
        'withExplanation': sum(1 for q in questions_json if q['explanationHtml']),
        'totalImages': sum(q['stemText'].count('<img ') + q['explanationHtml'].count('<img ') for q in questions_json),
        'totalTopics': count_topics(topics),
        'totalChapters': len(topics),
        'totalSubchapters': sum(len(t.get('children', [])) for t in topics),
    }

    errors = []
    for q in questions_json:
        qn_label = q['sequentialNumber']
        if q['answerMissing']:
            errors.append(f'Q{qn_label}: réponse manquante — désactivée')
        if len(q['propositions']) == 0:
            errors.append(f'Q{qn_label}: aucune proposition — désactivée')
        if 0 < len(q['propositions']) < 3:
            errors.append(f"Q{qn_label}: seulement {len(q['propositions'])} propositions — désactivée")
        if not q['answerMissing'] and len(q['answerLetters']) == 0 and len(q['propositions']) > 0:
            errors.append(f'Q{qn_label}: aucune réponse correcte — désactivée')

    return {'questions': questions_json, 'topics': topics, 'stats': result_stats, 'errors': errors}


# ---------------------------------------------------------------------------
# CLI entry point (kept for local dev)
# ---------------------------------------------------------------------------

def main():
    if len(sys.argv) < 2:
        print(json.dumps({'error': 'Usage: preprocessing_json.py <path-to-docx>'}))
        sys.exit(1)

    docx_path = sys.argv[1]
    if not Path(docx_path).exists():
        print(json.dumps({'error': f'File not found: {docx_path}'}))
        sys.exit(1)

    try:
        result = parse_docx_to_json(docx_path)
        json.dump(result, sys.stdout, ensure_ascii=False)
    except Exception as e:
        print(json.dumps({'error': str(e)}, ensure_ascii=False))
        sys.exit(1)


if __name__ == '__main__':
    main()
